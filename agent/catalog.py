"""Produktkatalog aus der Excel-Datei auf dem geteilten Laufwerk laden."""

import base64
import io

from openpyxl import load_workbook

from .graph import GraphClient


def _share_id(share_url: str) -> str:
    # Graph-Format für Freigabelinks: "u!" + base64url ohne Padding
    encoded = base64.urlsafe_b64encode(share_url.encode()).decode().rstrip("=")
    return f"u!{encoded}"


def download_excel(graph: GraphClient, share_url: str) -> bytes:
    return graph.request("GET", f"/shares/{_share_id(share_url)}/driveItem/content").content


def parse_products(xlsx: bytes, sheet: str | None = None) -> list[dict]:
    """Erste nicht-leere Zeile = Spaltenüberschriften, jede weitere Zeile = ein Produkt."""
    wb = load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = None
    for row in rows:
        if any(cell is not None for cell in row):
            header = [str(c).strip() if c is not None else f"Spalte_{i + 1}" for i, c in enumerate(row)]
            break
    if header is None:
        return []
    products = []
    for row in rows:
        if all(cell is None for cell in row):
            continue
        products.append({h: v for h, v in zip(header, row) if v is not None})
    return products


def load_catalog(graph: GraphClient, share_url: str, sheet: str | None = None) -> list[dict]:
    return parse_products(download_excel(graph, share_url), sheet)
